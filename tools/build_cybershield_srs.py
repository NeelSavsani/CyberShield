from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "CyberShield_SRS_Report.docx"

NAVY = "0B1E3D"
CYAN = "008BB5"
GRAY = "F2F4F7"
WHITE = "FFFFFF"
BLACK = "000000"
CONTENT_WIDTH_DXA = 9360


def set_font(run, size=15, bold=False, italic=False, color=BLACK):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.first_child_found_in("w:tblInd")
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_border(paragraph, color=NAVY):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepNext")
    p_pr.append(node)


def configure(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.45)
    sec.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(15)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in (
        ("Heading 1", 20, NAVY, 18, 10),
        ("Heading 2", 18, NAVY, 14, 7),
        ("Heading 3", 16, BLACK, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(15)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("CYBERSHIELD")
    set_font(header_run, 12, bold=True, color=NAVY)
    add_border(header, NAVY)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Page ")
    set_font(footer_run, 10, color=NAVY)
    add_page_field(footer)


def add_para(doc, text="", bold_label=None, italic=False, align=None, size=15, after=7):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if bold_label and text.startswith(bold_label):
        r = p.add_run(bold_label)
        set_font(r, size, bold=True)
        r = p.add_run(text[len(bold_label):])
        set_font(r, size, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, size, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(text)
    set_font(r, 15)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    r = p.add_run(text)
    set_font(r, 15)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, {1: 20, 2: 18, 3: 16}[level], bold=True, color=NAVY if level < 3 else BLACK)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, text in zip(hdr.cells, headers):
        shade(cell, GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, 12, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(text))
            set_font(r, 11)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def chapter_intro(doc, number, title, items):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(130)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(f"{number} {title.upper()}")
    set_font(r, 22, bold=True, color=NAVY)
    for item in items:
        add_bullet(doc, item)
    doc.add_page_break()


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run("CYBERSHIELD")
    set_font(r, 32, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Software Requirements Specification")
    set_font(r, 20, bold=True, color=CYAN)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(48)
    r = p.add_run("An Evidence-Driven URL Phishing Detection Platform")
    set_font(r, 17, italic=True, color=NAVY)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    values = [
        ("Document type", "Software Requirements Specification (SRS)"),
        ("Project status", "Implementation completed to date"),
        ("Prepared for", "Academic project documentation"),
        ("Date", date.today().strftime("%d %B %Y")),
    ]
    for row, (label, value) in zip(table.rows, values):
        shade(row.cells[0], GRAY)
        for index, cell in enumerate(row.cells):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(label if index == 0 else value)
            set_font(r, 13, bold=(index == 0), color=NAVY if index == 0 else BLACK)
    doc.add_page_break()


def front_matter(doc):
    add_heading(doc, "Acknowledgement", 1)
    add_para(doc, "This report records the requirements and current implementation of CyberShield, an academic cybersecurity project. It has been prepared from the implemented source code, user interface, analyzer pipeline, and supporting documentation available at the time of writing. The project acknowledges the guidance of faculty, peers, and all contributors who supported the research, development, testing, and documentation activities.")
    add_para(doc, "The document also recognizes the work of the open-source communities behind FastAPI, Playwright, Firebase, Python security libraries, and the web standards on which the prototype is built.")
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_para(doc, "CyberShield is a web-based URL phishing analysis platform. It investigates a submitted HTTP or HTTPS URL by collecting live infrastructure and page-behaviour evidence rather than relying only on URL character patterns or brand-name rules. The system validates the target, blocks localhost and private-network destinations, performs HTTP, DNS, domain, TLS, reputation, browser, HTML, JavaScript, and form analysis, then converts the evidence into a stable feature set for an explainable phishing-risk decision.")
    add_para(doc, "The current implementation provides a browser dashboard, Firebase-based user accounts and saved scan history, a FastAPI backend, an isolated Playwright browser session, screenshot capture, visible risk indicators, feature breakdowns, and printable analysis reports. The classifier is explicitly documented as an explainable development baseline; a separately generated synthetic-data model may support demonstrations but is not a production blocking model. This SRS defines the implemented scope, functional and non-functional requirements, data requirements, interfaces, constraints, and planned extensions.")
    doc.add_page_break()

    add_heading(doc, "Table of Contents", 1)
    toc = [
        "1 Introduction", "2 Overall Description", "3 System Features and Functional Requirements",
        "4 External Interface Requirements", "5 Non-functional Requirements", "6 Data Requirements and System Models",
        "7 Validation, Current Status and Future Scope", "8 References", "Appendix A - Requirement Traceability Matrix"
    ]
    for item in toc:
        add_para(doc, item, size=15, after=4)
    doc.add_page_break()

    add_heading(doc, "List of Figures", 1)
    for item in [
        "Figure 1. CyberShield high-level architecture",
        "Figure 2. URL analysis data flow",
        "Figure 3. User interaction and result lifecycle",
    ]:
        add_para(doc, item, size=15, after=4)
    add_heading(doc, "List of Tables", 1)
    for item in [
        "Table 1. Functional requirements", "Table 2. External interfaces", "Table 3. Quality attributes",
        "Table 4. Persistent data items", "Table 5. Requirement traceability matrix",
    ]:
        add_para(doc, item, size=15, after=4)


def chapter_one(doc):
    chapter_intro(doc, "1", "Introduction", ["Purpose and scope", "Objectives", "Problem definition", "Document conventions", "Literature and technology context"])
    add_heading(doc, "1.1 Purpose", 2)
    add_para(doc, "The purpose of CyberShield is to help a user inspect a potentially unsafe web address before interacting with it. The system presents an understandable phishing-risk result based on evidence observed from the public network, transport security, resolved page, browser behaviour, page markup, scripts, and forms. It is intended as decision support and not as an automatic authority to block or allow access.")
    add_heading(doc, "1.2 Scope", 2)
    add_para(doc, "The scope implemented to date is URL analysis. A user can submit a complete HTTP or HTTPS URL from the dashboard, view the returned analysis, retain it in personal history, revisit it from History or Reports, and print/download the browser result. Email and image tabs are present in the interface but are currently informational stubs and are not included in the implemented analysis scope.")
    add_heading(doc, "1.3 Objectives", 2)
    for text in [
        "Collect evidence from several independent sources so a decision is not based only on the URL text.",
        "Protect the analyzer from server-side request forgery by rejecting localhost, private, link-local, and other non-public destinations.",
        "Render the target in a controlled browser session and preserve a screenshot tied to that scan.",
        "Show a risk score, verdict, contributing indicators, and feature values in a user-friendly result page.",
        "Provide account-based persistence for scan history, reports, settings, and profile information.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "1.4 Problem Definition", 2)
    add_para(doc, "Phishing websites often imitate trusted services, redirect users across domains, collect credentials, request one-time passwords, use insecure form actions, or conceal risky behaviour in browser-rendered content. Users may see only a link or login page and cannot easily assess its technical properties. CyberShield addresses this gap by making collected risk evidence visible, structured, and reviewable before a user acts on a suspicious URL.")
    add_heading(doc, "1.5 Definitions and Abbreviations", 2)
    add_table(doc, ["Term", "Meaning"], [
        ("SRS", "Software Requirements Specification."),
        ("SSRF", "Server-Side Request Forgery; an attempt to make the analyzer access internal services."),
        ("TLS", "Transport Layer Security used to protect HTTPS connections."),
        ("DNS", "Domain Name System records used to resolve a hostname."),
        ("Firestore", "Firebase cloud document database used by the frontend for user-owned saved analyses."),
        ("Evidence-v1", "CyberShield's current stable feature schema derived from collector output."),
    ], [2500, 6860])


def chapter_two(doc):
    chapter_intro(doc, "2", "Overall Description", ["Product perspective", "User classes", "Operating environment", "Constraints and assumptions", "Technology stack"])
    add_heading(doc, "2.1 Product Perspective", 2)
    add_para(doc, "CyberShield is a multi-page web application with a separate API service. The frontend provides account and result workflows, while the backend coordinates live collectors. The application is designed as a demonstration-grade security analysis platform: it collects real evidence and clearly labels the current classifier limits.")
    add_heading(doc, "2.2 High-Level Architecture", 2)
    add_para(doc, "The browser sends a submitted URL to the FastAPI analysis endpoint. The orchestrator normalizes the URL, performs a public-target safety check, calls the collectors, derives evidence-v1 features, and requests a risk prediction. The frontend converts the response into a readable result and saves the user-visible record, including the screenshot URL and feature values, to the user's Firestore analysis collection.")
    add_table(doc, ["Layer", "Implemented responsibility", "Main technologies"], [
        ("Presentation", "Dashboard, history, reports, profile, settings, result rendering and browser print flow.", "HTML, CSS, JavaScript, Firebase SDK"),
        ("API", "URL validation, request handling, orchestration, CORS, health endpoint and static screenshot serving.", "FastAPI, Pydantic, Uvicorn"),
        ("Collection", "HTTP, DNS, domain, TLS, browser, HTML, JavaScript, form and reputation evidence collection.", "httpx, dnspython, WHOIS, Playwright, BeautifulSoup"),
        ("Decision", "Feature extraction, explainable contributors and optional demonstration model loading.", "Python, pandas, scikit-learn, joblib"),
        ("Persistence", "Authenticated user profiles, settings and personal analysis records; local API scan records support platform endpoints.", "Firebase Authentication, Firestore, SQLite"),
    ], [1500, 5000, 2860])
    add_heading(doc, "2.3 User Classes", 2)
    add_table(doc, ["User class", "Characteristics and permissions"], [
        ("Guest", "Can access the URL dashboard and use temporary session history. Profile, full History, and Reports are restricted."),
        ("Registered user", "Uses Firebase Authentication; can save and revisit personal analyses, update profile/settings, and access reports."),
        ("Administrator / developer", "Operates the local API, monitors diagnostics, configures optional threat-intelligence keys, and maintains the project."),
    ], [2700, 6660])
    add_heading(doc, "2.4 Operating Environment", 2)
    for text in [
        "Frontend: current desktop or mobile browser, with Chrome, Edge, or Firefox expected to support modern JavaScript and CSS.",
        "Backend: Python 3 environment capable of running FastAPI and Playwright Chromium. Windows is supported through a Proactor event loop for browser subprocess compatibility.",
        "Network: internet access is required for public DNS, HTTP/TLS, WHOIS, optional reputation providers, and Firebase cloud services.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "2.5 Constraints and Assumptions", 2)
    add_para(doc, "The analyzer only accepts complete public HTTP/HTTPS URLs. Live analysis quality depends on external websites, DNS resolution, TLS availability, browser compatibility, optional third-party keys, and network connectivity. A scan may return partial browser evidence when rendering fails; that failure must be shown as unavailable evidence rather than fabricated data. The user must understand that the current risk probability is advisory.")


def chapter_three(doc):
    chapter_intro(doc, "3", "System Features and Functional Requirements", ["Authentication and account workflow", "URL validation and analysis", "Risk decision and result presentation", "History and reports", "Administration and diagnostics"])
    add_heading(doc, "3.1 Functional Requirements", 2)
    reqs = [
        ("FR-01", "Account access", "The system shall allow registration, sign-in, password-reset flow, sign-out, and guest access through Firebase Authentication."),
        ("FR-02", "URL admission", "The system shall accept only complete HTTP/HTTPS URLs and reject malformed, localhost, private, link-local, reserved, or otherwise non-public targets before collection."),
        ("FR-03", "Infrastructure analysis", "The system shall collect HTTP reachability/redirect data, DNS data, domain age/registration data, and TLS certificate metadata when available."),
        ("FR-04", "Browser analysis", "The system shall open an admitted URL in a controlled Playwright browser context and collect the final URL, title, DOM, browser events, and a scan-specific screenshot when possible."),
        ("FR-05", "Content analysis", "The system shall inspect rendered HTML, JavaScript, and forms for observable phishing-relevant signals, including hidden iframes, obfuscated scripts, login form behaviour, OTP fields, card fields, insecure actions, and cross-domain form actions."),
        ("FR-06", "Risk result", "The system shall produce an explainable probability, risk class, verdict, and a ranked set of evidence contributors."),
        ("FR-07", "Result display", "The system shall display the risk score, verdict, analyzed input, timestamp, recommendation, indicators, feature breakdown, and screenshot when a screenshot exists for that scan."),
        ("FR-08", "Persistence", "The system shall save registered-user analyses with input type, content, score, verdict, indicators, features, screenshot URL, and timestamps in the user's Firestore collection."),
        ("FR-09", "History and reports", "The system shall list saved analyses and shall load the complete saved result, including features and screenshot URL, when a user selects View or View full result."),
        ("FR-10", "Reporting", "The system shall permit a saved result to be printed or saved as a PDF using the browser print facility."),
        ("FR-11", "Settings and profile", "The system shall allow an authenticated user to view/update supported profile fields and save report preferences."),
        ("FR-12", "Service health", "The API shall expose root and health endpoints that identify service status."),
    ]
    add_table(doc, ["ID", "Feature", "Requirement"], reqs, [900, 2050, 6410])
    add_heading(doc, "3.2 URL Analysis Workflow", 2)
    for text in [
        "Normalize the submitted address and validate the URL syntax.",
        "Resolve the hostname and reject any target that is local or non-public.",
        "Collect HTTP, DNS, domain, TLS, and optional reputation evidence.",
        "Render the page in the secure browser context, capture a screenshot, and inspect page content, scripts, and forms.",
        "Extract evidence-v1 features, calculate the phishing probability, identify contributors, and return the result.",
        "Persist the user-facing result and route the user to the complete analysis page.",
    ]:
        add_number(doc, text)
    add_heading(doc, "3.3 Risk Classification Rules", 2)
    add_para(doc, "The current baseline starts from a conservative base score and adds weighted evidence only when a collector supplies a value. Example evidence includes recent domain or certificate age, redirect chains, risky credential forms, OTP/card fields, hidden iframes, meta refresh redirects, downloads, permission prompts, obfuscated scripts, and configured reputation detections. The probability is mapped to Safe, Low, Medium, High, or Critical risk ranges. The frontend uses Safe below 40, Suspicious from 40 to 69, and Phishing at 70 or above for its user-facing verdict banner.")


def chapter_four(doc):
    chapter_intro(doc, "4", "External Interface Requirements", ["User interface", "Software interfaces", "API interfaces", "Communication interfaces"])
    add_heading(doc, "4.1 User Interface Requirements", 2)
    add_table(doc, ["Screen", "Required behaviour"], [
        ("Landing and authentication", "Provide a clear entry point, sign-in, registration, password recovery, and guest access."),
        ("Dashboard", "Provide a URL input, analysis progress states, current counters, recent history, and navigation to account features."),
        ("Analysis Result", "Show a readable verdict banner, risk score, input metadata, recommendation, indicators, feature table, scan-specific screenshot, and report actions."),
        ("History", "Provide filtering/searching, scan metadata, View full result, and deletion of the user's saved analysis."),
        ("Reports", "Show saved report rows and provide View access to the same complete result data as a newly completed scan."),
        ("Profile and Settings", "Allow account-related information and preferences to be viewed or changed according to authenticated access."),
    ], [2600, 6760])
    add_heading(doc, "4.2 API Interfaces", 2)
    add_table(doc, ["Endpoint", "Method", "Purpose"], [
        ("/analyze", "POST", "Runs synchronous URL analysis and returns the standardized analysis response."),
        ("/api/auth/register", "POST", "Creates a local platform user and signed token for the platform API workflow."),
        ("/api/auth/login", "POST", "Authenticates a local platform user and returns a signed bearer token."),
        ("/api/scans", "POST / GET", "Submits an asynchronous scan or retrieves scan history in the platform API workflow."),
        ("/api/scans/{scan_id}", "GET", "Returns the status/result of a queued scan after authorization checks."),
        ("/health", "GET", "Returns healthy/online service status."),
        ("/reports", "Static mount", "Serves generated analysis screenshot files for result rendering."),
    ], [2800, 1500, 5060])
    add_heading(doc, "4.3 Communication Interfaces", 2)
    add_para(doc, "The browser frontend communicates with the local FastAPI backend using HTTP requests and JSON. The backend communicates with public internet services over HTTP/HTTPS, DNS, TLS sockets, and WHOIS-related lookup mechanisms. Firebase Authentication and Firestore are accessed from the frontend through Firebase web modules. The deployment must configure CORS origins to allow only known frontend origins.")
    add_heading(doc, "4.4 Software Interfaces", 2)
    for text in [
        "FastAPI and Pydantic provide API routing and typed request/response models.",
        "Playwright Chromium renders target pages in a controlled browser context.",
        "BeautifulSoup and lxml support HTML parsing; httpx supports outbound HTTP.",
        "Firebase Authentication and Firestore support the production-facing frontend account and history experience.",
        "SQLite supports the implemented backend platform scan/account endpoints for local workflows.",
    ]:
        add_bullet(doc, text)


def chapter_five(doc):
    chapter_intro(doc, "5", "Non-functional Requirements", ["Security", "Performance and reliability", "Usability", "Maintainability", "Compliance and ethical use"])
    add_heading(doc, "5.1 Quality Attributes", 2)
    nfrs = [
        ("NFR-01", "Security", "Reject unsafe internal/private targets before HTTP, TLS, or browser collection to reduce SSRF exposure."),
        ("NFR-02", "Security", "Use Firebase Authentication for frontend session identity; enforce Firestore rules so users access only their records."),
        ("NFR-03", "Privacy", "Store only information needed for the scan record. A captured screenshot must be associated with the specific analysis that generated it."),
        ("NFR-04", "Reliability", "Return partial, labelled evidence when a collector is unavailable; do not convert missing evidence into a positive or negative claim."),
        ("NFR-05", "Performance", "Provide visible progress during analysis and avoid blocking the web server's primary event loop while Playwright creates browser subprocesses."),
        ("NFR-06", "Usability", "Present a concise verdict, explanation, risk score, and recommendation in language understandable to a non-specialist."),
        ("NFR-07", "Maintainability", "Keep collectors, feature extraction, classification, and frontend rendering separated so sources and weights can evolve independently."),
        ("NFR-08", "Auditability", "Preserve the input, timestamp, verdict, indicators, features, and screenshot reference for authenticated users' saved scans."),
        ("NFR-09", "Portability", "Support local development through a Python API and static frontend server, with a Docker Compose option for containerized setup."),
        ("NFR-10", "Ethical use", "Describe the current classifier as an advisory baseline. Do not represent a synthetic-data demonstration model as a validated production defence."),
    ]
    add_table(doc, ["ID", "Attribute", "Requirement"], nfrs, [900, 1800, 6660])
    add_heading(doc, "5.2 Security Requirements", 2)
    add_para(doc, "A security analyzer is itself exposed to hostile input. CyberShield shall treat every submitted URL as untrusted, block unsafe address classes before network access, use request timeouts, isolate webpage rendering from the user interface, and avoid exposing secret API keys in source control. Optional reputation providers must remain disabled until their keys are configured. User data access in Firestore must be scoped to the authenticated UID.")
    add_heading(doc, "5.3 Reliability and Error Handling", 2)
    add_para(doc, "If a target cannot be reached, a third-party lookup is unavailable, or browser analysis fails, the result shall preserve the error state and available evidence. The backend shall close browser resources in a finally path. The frontend shall hide a screenshot area when the requested image fails to load rather than displaying a wrong shared image.")


def chapter_six(doc):
    chapter_intro(doc, "6", "Data Requirements and System Models", ["Analysis response", "Feature schema", "Persistent data", "Data flow"])
    add_heading(doc, "6.1 Analysis Response Data", 2)
    add_para(doc, "The primary response contains success state, original and normalized URL, reachability, phishing probability, risk, message, and a data object. The data object holds source-specific evidence, the feature payload, the classifier output, and browser capture metadata.")
    add_table(doc, ["Data group", "Representative elements", "Purpose"], [
        ("Safety", "allowed, resolved_addresses", "Documents URL admission checks and resolved public addresses."),
        ("Infrastructure", "HTTP, DNS, domain, SSL/TLS, reputation", "Provides observable network and registration evidence."),
        ("Browser", "final URL, title, redirects, events, screenshot, browser error", "Records browser-rendered page behaviour and capture state."),
        ("Content", "HTML, JavaScript, form analysis", "Supplies page and credential-collection indicators."),
        ("Features", "schema_version, values, available_sources", "Produces stable model-ready evidence-v1 values."),
        ("Classification", "model version/status, probability, risk, contributors, notice", "Provides explainable user-facing decision information."),
    ], [1800, 4200, 3360])
    add_heading(doc, "6.2 Evidence-v1 Feature Families", 2)
    for text in [
        "Domain and TLS: domain age, certificate age, and derived recency risks.",
        "Network: DNSSEC availability, HTTP reachability, redirect count, and TLS availability.",
        "Forms: password, OTP, and payment-card fields; cross-domain and insecure form actions; GET login forms.",
        "Page and browser: hidden iframes, meta refresh, popups, downloads, permission requests, and JavaScript errors.",
        "Threat intelligence: detection count and Safe Browsing flag when providers are configured.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "6.3 Persistent User Analysis Record", 2)
    add_table(doc, ["Field", "Description", "Retention/use"], [
        ("inputType", "Type of user input; URL is currently supported.", "Used to render appropriate result metadata."),
        ("content", "Submitted URL.", "Shown in history and result metadata."),
        ("riskScore and verdict", "User-facing assessment outputs.", "Used for lists, statistics, filters, and result banner."),
        ("indicators and features", "Contributors and stable feature values.", "Used to recreate the complete result detail."),
        ("screenshotUrl", "URL for the capture generated by this scan.", "Used only for the matching result view."),
        ("analyzedAt and createdAt", "Scan and persistence timestamps.", "Used for ordering, display, and reporting."),
    ], [2200, 4200, 2960])
    add_heading(doc, "6.4 Data Flow", 2)
    add_para(doc, "User -> Dashboard -> POST /analyze -> URL safety gate -> evidence collectors -> feature extractor -> explainable classifier -> normalized frontend result -> Firestore user analysis record -> Result / History / Reports. A selected saved record is reconstructed with the same indicators, features, and screenshot URL before the Result page renders it.")


def chapter_seven(doc):
    chapter_intro(doc, "7", "Validation, Current Status and Future Scope", ["Validation approach", "Implemented status", "Known limitations", "Future enhancements"])
    add_heading(doc, "7.1 Requirement Validation Approach", 2)
    add_para(doc, "Validation is performed through source-level checks, API execution, browser workflow inspection, and result-data consistency checks. The project includes decision-pipeline tests for safety and classification behaviour. Manual acceptance checks confirm registration/guest paths, a successful public URL scan, result rendering, screenshot display, history filtering, report opening, and deletion behaviour.")
    add_heading(doc, "7.2 Implemented Status", 2)
    add_table(doc, ["Area", "Status", "Notes"], [
        ("Public URL safety gate", "Implemented", "Blocks localhost and non-public IP destinations before collectors run."),
        ("Live infrastructure collection", "Implemented", "HTTP, DNS, domain, TLS, and optional reputation evidence are collected where available."),
        ("Controlled browser collection", "Implemented", "Playwright captures rendered evidence and a scan-specific screenshot."),
        ("Explainable risk output", "Implemented", "Evidence baseline is active; an optional synthetic-data demo model is not production validated."),
        ("Frontend account/history/report workflow", "Implemented", "Firebase account and Firestore persistence are used by the frontend."),
        ("Complete saved result reconstruction", "Implemented", "History and Reports pass indicators, features, and the screenshot URL to the Result page."),
        ("Email and image analysis", "Planned", "Interface tabs are present but currently show an availability notice."),
        ("OpenPhish and URLhaus feeds", "Planned", "Provider placeholders are present; integration is pending."),
    ], [2850, 1400, 5110])
    add_heading(doc, "7.3 Known Limitations", 2)
    for text in [
        "A risk score is not a guarantee of safety or maliciousness; human verification remains necessary.",
        "The baseline classifier must be trained and evaluated on labelled benign and phishing samples before automated blocking is considered.",
        "Website rendering and external lookups can fail because of target-side restrictions, network conditions, or provider configuration.",
        "Screenshot availability is dependent on successful browser capture and continued availability of the served capture file.",
        "The frontend currently uses Firebase identity/history while the backend also contains a separate local platform account/scan API path; deployment should select and consolidate the intended production identity strategy.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "7.4 Future Scope", 2)
    for text in [
        "Integrate and poll VirusTotal, Google Safe Browsing, OpenPhish, and URLhaus according to their service policies.",
        "Collect consented, labelled evidence rows; train versioned models; measure precision, recall, false-positive rate, and calibration.",
        "Add email/text and uploaded-image analysis with equivalent persistence and explanation quality.",
        "Introduce background jobs, queue monitoring, retry policy, retention controls, and production observability.",
        "Harden deployment with secret management, HTTPS, restrictive CORS, Firebase Security Rules review, rate limiting, and security logging.",
    ]:
        add_bullet(doc, text)


def references_and_appendix(doc):
    doc.add_page_break()
    add_heading(doc, "8 References", 1)
    refs = [
        "I. Sommerville, Software Engineering, Pearson. Guidance for requirements engineering and SRS structure.",
        "ISO/IEC/IEEE 29148, Systems and software engineering - Life cycle processes - Requirements engineering.",
        "FastAPI Documentation, https://fastapi.tiangolo.com/.",
        "Playwright Python Documentation, https://playwright.dev/python/.",
        "Firebase Documentation, https://firebase.google.com/docs/.",
        "OWASP Server-Side Request Forgery Prevention Cheat Sheet, https://cheatsheetseries.owasp.org/.",
        "Google Safe Browsing API Documentation, https://developers.google.com/safe-browsing/.",
        "CyberShield project source code and README, current workspace implementation, accessed on the report date.",
    ]
    for ref in refs:
        add_bullet(doc, ref)
    doc.add_page_break()
    add_heading(doc, "Appendix A - Requirement Traceability Matrix", 1)
    add_table(doc, ["Requirement", "Implementation evidence", "Validation method"], [
        ("FR-02", "url_safety.py and orchestrator admission path", "Submit malformed, localhost, and private-address URLs; verify denial before collection."),
        ("FR-03 to FR-05", "Analyzer modules and orchestrator collector sequence", "Run an admitted public URL and inspect returned source evidence."),
        ("FR-06", "FeatureExtractor and EvidenceClassifier", "Verify feature values, probability, risk, and contributor list in the response."),
        ("FR-07", "result.html and result.js", "Confirm score, indicators, features, recommendation, metadata, and matching screenshot render."),
        ("FR-08 to FR-09", "firebase-data.js, history.js, reports.js", "Save a result, reopen it from History and Reports, and compare full detail with fresh result."),
        ("NFR-04", "Browser exception handling and available_sources", "Simulate/observe a browser failure and verify partial evidence is labelled."),
        ("NFR-10", "README and classifier notice", "Review that baseline/demo model limits are presented to users/developers."),
    ], [1500, 4600, 3260])
    add_heading(doc, "Appendix B - Formatting Compliance", 1)
    add_para(doc, "This report follows the visual conventions observed in the supplied Expirochain report: A4 portrait pages, one-inch margins, Arial typography, large readable body text, bold numbered chapter and section headings, a running project header, centered page numbering, consistent table padding, and controlled white space. The SRS content itself follows a requirements-oriented structure so that implemented functionality, current boundaries, and future work remain distinguishable.")


def main():
    doc = Document()
    configure(doc)
    cover(doc)
    front_matter(doc)
    chapter_one(doc)
    chapter_two(doc)
    chapter_three(doc)
    chapter_four(doc)
    chapter_five(doc)
    chapter_six(doc)
    chapter_seven(doc)
    references_and_appendix(doc)
    doc.core_properties.title = "CyberShield Software Requirements Specification"
    doc.core_properties.subject = "Academic project SRS"
    doc.core_properties.author = "CyberShield Project Team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path('output'); OUT.mkdir(exist_ok=True); DEST=OUT/'CyberShield_Team_Explainer.docx'
NAVY=RGBColor(11,30,61); CYAN=RGBColor(0,120,170); GRAY=RGBColor(80,80,80)

def font(run,size=11,bold=False,color=None,name='Calibri'):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name); run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=color
def shade(cell,fill='E8EEF5'):
    p=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); p.append(e)
def margins(cell):
    p=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:tcMar'); p.append(e)
    for s in ['top','start','bottom','end']:
        n=OxmlElement('w:'+s); n.set(qn('w:w'),'100'); n.set(qn('w:type'),'dxa'); e.append(n)
def add_page_number(p):
    r=p.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.append(f)

d=Document(); sec=d.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=Inches(0.85); sec.left_margin=sec.right_margin=Inches(0.9)
styles=d.styles; n=styles['Normal']; n.font.name='Calibri'; n._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); n.font.size=Pt(11); n.paragraph_format.space_after=Pt(6); n.paragraph_format.line_spacing=1.1
for name,size,col in [('Heading 1',16,NAVY),('Heading 2',13,NAVY),('Heading 3',11,RGBColor(31,77,120))]:
    s=styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=col; s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6)
head=sec.header.paragraphs[0]; head.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=head.add_run('CYBERSHIELD | TEAM EXPLAINER'); font(r,8,True,GRAY)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=foot.add_run('CyberShield Team Guide | '); font(r,9,False,GRAY); add_page_number(foot)

def p(text='',align=None,bold_prefix=None,after=6):
    x=d.add_paragraph(); x.paragraph_format.space_after=Pt(after); x.paragraph_format.line_spacing=1.1
    if align is not None:x.alignment=align
    if bold_prefix and text.startswith(bold_prefix):
        r=x.add_run(bold_prefix); font(r,11,True); r=x.add_run(text[len(bold_prefix):]); font(r)
    else: r=x.add_run(text); font(r)
    return x
def h(text,l=1): return d.add_paragraph(text,style=f'Heading {l}')
def b(text):
    x=d.add_paragraph(style='List Bullet'); x.paragraph_format.space_after=Pt(3); x.paragraph_format.line_spacing=1.1; r=x.add_run(text); font(r); return x
def num(text):
    x=d.add_paragraph(style='List Number'); x.paragraph_format.space_after=Pt(3); x.paragraph_format.line_spacing=1.1; r=x.add_run(text); font(r); return x
def tbl(caption,headers,rows,widths=None):
    c=d.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(caption); font(r,10,True,NAVY)
    t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for cell,val in zip(t.rows[0].cells,headers):
        shade(cell); margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; r=cell.paragraphs[0].add_run(val); font(r,9,True)
    for row in rows:
        for cell,val in zip(t.add_row().cells,row):
            margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; r=cell.paragraphs[0].add_run(val); font(r,9)
    if widths:
        t.autofit=False
        for row in t.rows:
            for cell,w in zip(row.cells,widths):cell.width=Inches(w)
    d.add_paragraph()
def callout(label,text):
    t=d.add_table(rows=1,cols=1); t.style='Table Grid'; cell=t.cell(0,0); shade(cell,'EEF7FB'); margins(cell)
    x=cell.paragraphs[0]; r=x.add_run(label+' '); font(r,10,True,NAVY); r=x.add_run(text); font(r,10)
    d.add_paragraph()
def code(lines):
    t=d.add_table(rows=1,cols=1); t.style='Table Grid'; cell=t.cell(0,0); shade(cell,'F4F6F8'); margins(cell)
    x=cell.paragraphs[0]; x.paragraph_format.line_spacing=1.0
    for i,line in enumerate(lines):
        r=x.add_run(line+('\n' if i<len(lines)-1 else '')); font(r,9,False,None,'Courier New')
    d.add_paragraph()

# Cover
for _ in range(5):p('',after=10)
x=p('CYBERSHIELD',WD_ALIGN_PARAGRAPH.CENTER,after=6); font(x.runs[0],28,True,CYAN)
x=p('Team Explanation and Faculty Viva Guide',WD_ALIGN_PARAGRAPH.CENTER,after=18); font(x.runs[0],18,True,NAVY)
p('How the system works, the technologies we used, how to demonstrate it, and how to answer likely questions.',WD_ALIGN_PARAGRAPH.CENTER,after=36)
p('For: CyberShield Project Team',WD_ALIGN_PARAGRAPH.CENTER,after=5)
p('Prepared from the current project implementation',WD_ALIGN_PARAGRAPH.CENTER,after=5)
p('Replace this line with team names before printing or sharing.',WD_ALIGN_PARAGRAPH.CENTER,after=0)
d.add_page_break()

h('1. One-Minute Explanation',1)
callout('Say this first:', 'CyberShield is a web-based phishing URL analysis platform. A user submits a website URL, and our backend safely investigates the site using live evidence such as redirects, DNS, domain age, TLS certificate details, page forms, browser behaviour, and selected threat-intelligence signals. We combine this evidence into an explainable risk score and show why the URL was classified as Safe, Low, Medium, High, or Critical.')
p('The important difference is that CyberShield is not based only on simple URL rules such as length or suspicious words. It actually collects technical signals from the target website and its infrastructure. This makes the result more meaningful and easier to explain to a user.')
h('What problem are we solving?',2)
b('Users often receive suspicious links through email, messages, social media, or fake login pages.')
b('Many people cannot judge a URL by looking at it, especially when attackers use redirects, valid certificates, and convincing page designs.')
b('CyberShield gives a structured risk assessment before the user enters credentials or payment data.')
h('2. The System in One Flow',1)
code(['User enters URL in dashboard','        |','Normalize and validate URL','        |','Block localhost / private IP / unsafe target','        |','Collect HTTP, DNS, domain, TLS, browser, HTML, form and reputation evidence','        |','Extract stable evidence-v1 features','        |','Run explainable baseline classifier','        |','Display probability, risk label, contributors and technical details'])
callout('Key message:', 'The safety check happens before network collection. This prevents the tool from being used to access private or local systems through a submitted URL.')

h('3. Step-by-Step: How CyberShield Works',1)
h('3.1 URL input, normalization and validation',2)
p('The user enters a web address in the dashboard. The backend normalizes the value and verifies that it is a complete HTTP or HTTPS URL. Unsupported schemes and malformed inputs are rejected with a clear message. This is our first reliability and security gate.')
h('3.2 SSRF protection: blocking unsafe targets',2)
p('A URL scanner is itself a network-facing tool, so it must not blindly fetch any address a user supplies. CyberShield blocks localhost, names ending in .localhost, private IP addresses, loopback addresses, link-local addresses, and other non-public destinations. It resolves hostnames and checks their addresses before allowing analysis.')
callout('Good viva answer:', 'This control protects against Server-Side Request Forgery, or SSRF. Without it, an attacker could try to use our backend to reach internal services that are not exposed to the public internet.')
h('3.3 HTTP and redirect evidence',2)
p('The HTTP analyzer checks whether the URL is reachable and records response behaviour. Redirect chains are relevant because phishing pages may use several redirections to hide the actual destination. If a URL is not reachable, CyberShield returns Unknown rather than incorrectly declaring it safe.')
h('3.4 DNS, domain and TLS evidence',2)
p('The DNS analyzer observes domain resolution. The domain analyzer obtains registration information such as domain age when available. Very recent domains can be a useful risk signal, but they are not proof by themselves. The TLS analyzer gathers certificate metadata. A valid HTTPS certificate helps encrypt traffic, but it does not prove that a website is legitimate; phishing sites can also use HTTPS.')
h('3.5 Controlled browser analysis',2)
p('CyberShield uses Playwright with Chromium to load the page in a controlled browser session. This is important because many websites build content with JavaScript after the first HTTP response. The browser session captures the final URL, title, HTML size, redirects, cookie count, popups, downloads, permission requests, JavaScript errors, network counts, and a screenshot where possible.')
h('3.6 HTML, JavaScript and form analysis',2)
p('After browser loading, separate analyzers inspect the rendered page. The HTML analyzer identifies structural signals such as hidden iframes and meta-refresh behaviour. The JavaScript analyzer looks for suspicious script patterns such as obfuscation. The form analyzer checks whether a login or credential form sends data to a different domain, uses insecure HTTP, uses GET for password submission, or requests OTP and payment-card data.')
h('3.7 Feature extraction and explainable decision',2)
p('The feature extractor converts raw analyzer outputs into a stable evidence-v1 feature payload. The classifier starts with a base score and adds weighted impact for evidence that is present. It converts the score into a probability and maps that probability to a risk label. It also returns the highest-impact contributors, so the user sees reasons such as a cross-domain credential form, a recently registered domain, or a threat-intelligence detection.')

h('4. Technologies We Used and Why',1)
tbl('Technology Stack', ['Technology','Where we use it','Why it is useful'], [
['HTML, CSS, JavaScript','Frontend pages and dashboard','Creates a responsive, multi-page web interface.'],
['Python','Backend services and analysis logic','Clear, modular programming for data collection and APIs.'],
['FastAPI','REST API backend','Fast request handling, validation, and automatic API documentation.'],
['Pydantic','Request/response models','Ensures structured and validated API data.'],
['Playwright + Chromium','Browser analysis','Loads JavaScript-rendered pages in a controlled session.'],
['httpx / requests','HTTP collection','Collects page and redirect information.'],
['dnspython','DNS analysis','Resolves and inspects domain information.'],
['python-whois','Domain age information','Helps identify recently registered domains where data exists.'],
['cryptography / pyOpenSSL','TLS inspection','Reads certificate-related security metadata.'],
['BeautifulSoup + lxml','HTML parsing','Extracts page structure and form signals.'],
['SQLite','Development data storage','Stores local application data and history simply.'],
['scikit-learn, pandas, joblib','Demo ML workflow','Supports versioned training experiments and model loading.'],
['Docker Compose','Optional deployment','Makes local setup more repeatable.']], [1.55,2.1,2.75])
h('What must we not claim?',2)
p('We must not say that a valid TLS certificate means a website is safe. We must not say that the system guarantees detection of every phishing website. We must not say that the current baseline classifier is a production model trained on large real-world phishing data. The repository clearly states that the current development model uses synthetic data for a college demonstration. Our strongest answer is that CyberShield is an explainable evidence-collection pipeline with an honest baseline decision layer.')

h('5. Architecture and Module Responsibilities',1)
tbl('Main Modules', ['Module','Responsibility','Output'], [
['Frontend','Collect URL and display result, history, reports, accounts','User-friendly pages and requests'],
['FastAPI API','Receives request and sends structured response','JSON API response'],
['URL safety service','Rejects unsafe local/private targets','Allowed status or safe rejection'],
['Orchestrator','Calls analyzers in the correct order and merges data','Analysis data object'],
['Analyzers','Collect specialized evidence independently','HTTP, DNS, TLS, browser, form and other signals'],
['Feature extractor','Converts raw evidence into evidence-v1 values','Model-ready features'],
['Evidence classifier','Computes probability, risk and contributors','Explainable classification'],
['Database/history','Preserves analysis records where enabled','Reviewable stored results']], [1.55,3.05,1.8])
h('Why use an orchestrator?',2)
p('The orchestrator is the central coordinator. It keeps the analysis process organized: normalize URL, validate it, apply the safety check, run evidence collectors, handle browser errors, extract features, classify, and return a standard response. This design is easier to test and maintain than putting every operation inside one large API route.')
h('Why use multiple analyzers?',2)
p('Different evidence sources change independently. For example, a browser failure should not necessarily stop DNS or TLS collection. Keeping analyzers separate makes the system modular, makes fault handling clearer, and allows future improvements such as adding a visual brand-matching analyzer or new threat-intelligence provider.')

h('6. How to Demonstrate the Project',1)
h('Recommended live demo order',2)
num('Open the CyberShield home page and state the problem: users cannot easily judge suspicious links.')
num('Open the dashboard and submit a normal public HTTPS URL. Explain that CyberShield first validates the URL and blocks private targets.')
num('Show the result screen. Point out the risk label, phishing probability, evidence contributors, technical details, and screenshot if available.')
num('Open history or reports to show that results can be reviewed later.')
num('Try localhost or a private address only if the live UI/API supports the demonstration safely. Explain that rejection is an SSRF protection feature, not a phishing verdict.')
num('Conclude with the limitation: the present classifier is a baseline demo; production use requires a validated model trained on real labelled data.')
callout('Presenter tip:', 'During the demonstration, do not visit a real dangerous phishing page just to prove the system. Use safe public examples or controlled test pages. A responsible demonstration is a strength, not a weakness.')

h('7. Suggested Team Roles During Viva',1)
tbl('Speaking Plan', ['Team member focus','What to explain'], [
['Member 1: Problem and UI','Problem statement, target users, dashboard flow, result readability.'],
['Member 2: Backend flow','FastAPI API, request lifecycle, orchestrator, JSON response.'],
['Member 3: Security','SSRF protection, public-IP checks, controlled browser, non-submission of forms.'],
['Member 4: Detection logic','Evidence sources, feature extraction, baseline scoring, contributors and limitations.'],
['Any member: Future scope','Real labelled dataset, model metrics, reputation APIs, queueing and deployment hardening.']], [2.0,4.4])
h('Useful handoff sentence',2)
p('After explaining your part, say: “This is how the [module] works. Now [team member] will explain how its output is used in the next stage of the pipeline.” This makes the presentation sound coordinated.')

h('8. Likely Mentor and Faculty Questions',1)
qa=[
('What is the main objective of CyberShield?', 'To help a user assess a submitted web URL using live, explainable evidence before interacting with a potentially deceptive page.'),
('Why is this better than checking URL length or keywords?', 'Those rules are easy to bypass and provide weak explanations. CyberShield combines infrastructure, redirect, browser, DOM, form, and optional reputation signals.'),
('What is SSRF and how do you prevent it?', 'SSRF is when a server is tricked into making requests to internal resources. We block localhost, private, loopback, link-local, and non-public resolved IP addresses before analysis.'),
('Why do you use Playwright?', 'Many modern phishing pages are JavaScript-rendered. Playwright lets us observe the rendered DOM and browser behaviour in a controlled browser context.'),
('Do you submit forms or enter credentials during analysis?', 'No. We observe form structure and action targets only. We do not submit target-site forms or enter user credentials.'),
('Does HTTPS mean the URL is safe?', 'No. HTTPS only secures communication between the browser and site. A phishing site can also have a valid certificate, so we treat TLS as one signal among many.'),
('How does the risk score work?', 'The evidence baseline starts from a base score and adds weighted impact for observable risk signals. It returns a probability, a label, and the most influential contributors.'),
('Is the ML model production ready?', 'No. The current model is explicitly a college-demo baseline trained from synthetic data. Production blocking requires lawful labelled data, validation, versioning, and measured precision/recall.'),
('What happens if a website cannot be reached?', 'CyberShield returns Unknown or an error state. It does not convert missing evidence into a Safe result.'),
('What are the limitations?', 'External data can be unavailable, websites can block automation, and a baseline model can have false positives and false negatives. The tool is decision support, not a guarantee.'),
('Why store history?', 'History helps users revisit analyses, supports transparency, and provides a foundation for later reporting and evaluation.'),
('What can you improve next?', 'Add opt-in reputation providers, collect labelled datasets, evaluate model metrics, use async queues for long scans, improve authorization, and add data retention controls.')]
for q,a in qa:
    p('Q: '+q, bold_prefix='Q: '); p('Answer: '+a, bold_prefix='Answer: ',after=8)

h('9. Strong Closing Statement',1)
callout('Closing statement:', 'CyberShield is not just a URL keyword checker. It is a modular, security-conscious analysis system that safely collects live web evidence, explains its risk decision, and clearly communicates the current model limitations. Our project demonstrates a practical foundation for an accountable phishing-detection platform.')
h('10. Team Checklist Before Meeting Faculty',1)
for item in ['Each member can explain the full pipeline in simple language.', 'Each member knows why private IP addresses and localhost are blocked.', 'Each member can explain why HTTPS is not proof of legitimacy.', 'Each member knows the difference between evidence collection and classification.', 'Each member states the baseline-model limitation honestly.', 'The team uses only safe URLs or controlled pages during a demo.', 'Team members know who will explain frontend, backend, security, and model logic.']:
    b(item)

d.core_properties.title='CyberShield Team Explanation and Faculty Viva Guide'; d.core_properties.author='CyberShield Team'; d.save(DEST); print(DEST)

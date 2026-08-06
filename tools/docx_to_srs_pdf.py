from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "output" / "CyberShield_SRS_Report.docx"
OUTPUT = ROOT / "output" / "CyberShield_SRS_Report.pdf"
FONT_DIR = Path(r"C:\Windows\Fonts")
FONT = "Arial"
FONT_BOLD = "Arial-Bold"


def register_fonts():
    pdfmetrics.registerFont(TTFont(FONT, str(FONT_DIR / "arial.ttf")))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, str(FONT_DIR / "arialbd.ttf")))


def styles():
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("body", parent=base["BodyText"], fontName=FONT, fontSize=13.8, leading=18, alignment=TA_JUSTIFY, spaceAfter=8),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName=FONT_BOLD, fontSize=18, leading=23, textColor=colors.HexColor("#0B1E3D"), spaceBefore=18, spaceAfter=10, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName=FONT_BOLD, fontSize=15.5, leading=19, textColor=colors.HexColor("#0B1E3D"), spaceBefore=14, spaceAfter=7, keepWithNext=True),
        "h3": ParagraphStyle("h3", parent=base["Heading3"], fontName=FONT_BOLD, fontSize=13.5, leading=17, textColor=colors.black, spaceBefore=10, spaceAfter=5, keepWithNext=True),
        "cover": ParagraphStyle("cover", parent=base["Title"], fontName=FONT_BOLD, fontSize=27, leading=33, alignment=TA_CENTER, textColor=colors.HexColor("#0B1E3D"), spaceAfter=18),
        "cover_sub": ParagraphStyle("cover_sub", parent=base["BodyText"], fontName=FONT_BOLD, fontSize=17, leading=22, alignment=TA_CENTER, textColor=colors.HexColor("#008BB5"), spaceAfter=22),
        "cover_italic": ParagraphStyle("cover_italic", parent=base["BodyText"], fontName=FONT, fontSize=14, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#0B1E3D"), spaceAfter=42),
        "bullet": ParagraphStyle("bullet", parent=base["BodyText"], fontName=FONT, fontSize=13.8, leading=18, leftIndent=20, firstLineIndent=-10, bulletIndent=6, spaceAfter=5, alignment=TA_LEFT),
        "table": ParagraphStyle("table", parent=base["BodyText"], fontName=FONT, fontSize=10.2, leading=12.2, alignment=TA_LEFT),
        "table_hdr": ParagraphStyle("table_hdr", parent=base["BodyText"], fontName=FONT_BOLD, fontSize=10.4, leading=12.4, alignment=TA_CENTER, textColor=colors.HexColor("#0B1E3D")),
    }


def text_of(paragraph):
    return paragraph.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def is_page_break(paragraph):
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def paragraph_story(paragraph, s):
    text = text_of(paragraph)
    if is_page_break(paragraph):
        return [PageBreak()]
    if not text.strip():
        return [Spacer(1, 5)]
    style_name = paragraph.style.name
    if style_name == "Heading 1":
        return [Paragraph(text, s["h1"])]
    if style_name == "Heading 2":
        return [Paragraph(text, s["h2"])]
    if style_name == "Heading 3":
        return [Paragraph(text, s["h3"])]
    if style_name in {"List Bullet", "List Number"}:
        bullet = "•" if style_name == "List Bullet" else "•"
        return [Paragraph(text, s["bullet"], bulletText=bullet)]
    if text == "CYBERSHIELD":
        return [Spacer(1, 1.3 * inch), Paragraph(text, s["cover"])]
    if text == "Software Requirements Specification":
        return [Paragraph(text, s["cover_sub"])]
    if text == "An Evidence-Driven URL Phishing Detection Platform":
        return [Paragraph(text, s["cover_italic"])]
    return [Paragraph(text, s["body"])]


def table_story(table, s, width):
    rows = []
    for ri, row in enumerate(table.rows):
        cells = []
        for ci, cell in enumerate(row.cells):
            style = s["table_hdr"] if ri == 0 else s["table"]
            cells.append(Paragraph(cell.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), style))
        rows.append(cells)
    n = len(rows[0])
    if n == 2:
        col_widths = [width * .29, width * .71]
    elif n == 3:
        col_widths = [width * .15, width * .27, width * .58]
    else:
        col_widths = [width / n] * n
    out = Table(rows, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#AAB7C4")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if n == 2 and len(rows) == 4 and rows[0][0].getPlainText() == "Document type":
        commands.extend([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F7")), ("BACKGROUND", (1, 0), (1, -1), colors.white)])
    out.setStyle(TableStyle(commands))
    return [out, Spacer(1, 9)]


def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", next(p for p in parent.paragraphs if p._p == child)
        elif child.tag == qn("w:tbl"):
            yield "t", next(t for t in parent.tables if t._tbl == child)


def header_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#0B1E3D"))
    canvas.setLineWidth(.75)
    canvas.line(doc.leftMargin, A4[1] - .58 * inch, A4[0] - doc.rightMargin, A4[1] - .58 * inch)
    canvas.setFont(FONT_BOLD, 8.5)
    canvas.setFillColor(colors.HexColor("#0B1E3D"))
    canvas.drawString(doc.leftMargin, A4[1] - .45 * inch, "CYBERSHIELD")
    canvas.setFont(FONT, 8.5)
    canvas.drawCentredString(A4[0] / 2, .42 * inch, f"Page {doc.page}")
    canvas.restoreState()


def main():
    register_fonts()
    source = Document(INPUT)
    s = styles()
    doc = SimpleDocTemplate(str(OUTPUT), pagesize=A4, leftMargin=inch, rightMargin=inch, topMargin=.85 * inch, bottomMargin=.75 * inch, title="CyberShield Software Requirements Specification", author="CyberShield Project Team")
    story = []
    width = A4[0] - doc.leftMargin - doc.rightMargin
    for kind, block in iter_blocks(source):
        story.extend(paragraph_story(block, s) if kind == "p" else table_story(block, s, width))
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    print(OUTPUT)


if __name__ == "__main__":
    main()

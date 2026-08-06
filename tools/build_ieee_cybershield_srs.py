from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
DIAGRAMS = OUT / "ieee_srs_diagrams"
OUT.mkdir(exist_ok=True)
DIAGRAMS.mkdir(exist_ok=True)
OUTPUT = OUT / "CyberShield_IEEE_SRS.docx"
BLACK = "000000"
GRAY = "E7E6E6"
TABLE_WIDTH = 9000


def font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(BLACK)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(cell, value=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), value)
    tc_pr.append(shd)


def table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(.45)
    sec.footer_distance = Inches(.45)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, italic in (("Heading 1", 16, False), ("Heading 2", 14, False), ("Heading 3", 13, True)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = italic
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run("CYBERSHIELD - SOFTWARE REQUIREMENTS SPECIFICATION")
    font(r, 10, bold=True)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field(footer)


def p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False, italic=False, before=0, after=0):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    r = para.add_run(text)
    font(r, size, bold, italic)
    return para


def heading(doc, text, level):
    para = doc.add_paragraph(style=f"Heading {level}")
    r = para.add_run(text)
    font(r, {1: 16, 2: 14, 3: 13}[level], bold=True, italic=(level == 3))
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.left_indent = Inches(.35)
    para.paragraph_format.first_line_indent = Inches(-.18)
    r = para.add_run(text)
    font(r, 12)
    return para


def number(doc, text):
    para = doc.add_paragraph(style="List Number")
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.left_indent = Inches(.38)
    para.paragraph_format.first_line_indent = Inches(-.20)
    r = para.add_run(text)
    font(r, 12)
    return para


def caption(doc, text, above=True):
    para = p(doc, text, WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True, before=5 if above else 3, after=3)
    return para


def table(doc, headers, rows, widths, caption_text=None):
    if caption_text:
        caption(doc, caption_text, above=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    table_geometry(t, widths)
    for c, text in zip(t.rows[0].cells, headers):
        shade(c)
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.line_spacing = 1
        r = para.add_run(text); font(r, 10.5, bold=True)
    for values in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, values):
            para = c.paragraphs[0]; para.paragraph_format.line_spacing = 1
            r = para.add_run(str(text)); font(r, 10)
    table_geometry(t, widths)
    p(doc, "", size=3, after=2)
    return t


def diagram_canvas(path, title, boxes, arrows):
    w, h = 1400, 860
    im = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(im)
    try:
        f_title = ImageFont.truetype("C:/Windows/Fonts/timesbd.ttf", 34)
        f_box = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 24)
    except OSError:
        f_title = f_box = ImageFont.load_default()
    draw.text((w // 2, 28), title, fill="black", font=f_title, anchor="ma")
    for x1, y1, x2, y2 in arrows:
        draw.line((x1, y1, x2, y2), fill="black", width=3)
        dx, dy = x2 - x1, y2 - y1
        if abs(dx) >= abs(dy):
            sign = 1 if dx >= 0 else -1
            draw.polygon([(x2, y2), (x2 - 15 * sign, y2 - 9), (x2 - 15 * sign, y2 + 9)], fill="black")
        else:
            sign = 1 if dy >= 0 else -1
            draw.polygon([(x2, y2), (x2 - 9, y2 - 15 * sign), (x2 + 9, y2 - 15 * sign)], fill="black")
    for x, y, bw, bh, label in boxes:
        draw.rounded_rectangle((x, y, x + bw, y + bh), radius=12, outline="black", width=3, fill="white")
        lines = label.split("\n")
        cy = y + bh / 2 - (len(lines) - 1) * 16
        for line in lines:
            draw.text((x + bw / 2, cy), line, fill="black", font=f_box, anchor="mm")
            cy += 32
    im.save(path)


def make_diagrams():
    use_case = DIAGRAMS / "use_case.png"
    diagram_canvas(use_case, "CyberShield Use Case Model", [
        (65, 330, 220, 85, "Guest User"), (65, 520, 220, 85, "Registered User"),
        (565, 120, 290, 85, "Submit URL"), (565, 275, 290, 85, "View Analysis Result"),
        (565, 430, 290, 85, "View History / Reports"), (565, 585, 290, 85, "Manage Profile / Settings"),
        (1060, 350, 230, 85, "Administrator"),
    ], [(285, 372, 565, 162), (285, 372, 565, 317), (285, 562, 565, 317), (285, 562, 565, 472), (285, 562, 565, 627), (855, 317, 1060, 392)])
    dfd = DIAGRAMS / "dfd.png"
    diagram_canvas(dfd, "CyberShield Level-0 Data Flow Diagram", [
        (55, 355, 220, 85, "User"), (420, 320, 300, 110, "CyberShield\nAnalysis Platform"),
        (930, 145, 300, 85, "Public Web / DNS / TLS"), (930, 345, 300, 85, "Optional Reputation APIs"),
        (930, 545, 300, 85, "Firebase / Firestore"),
    ], [(275, 398, 420, 375), (720, 355, 930, 188), (720, 375, 930, 388), (720, 398, 930, 588), (930, 230, 720, 355), (930, 430, 720, 375), (930, 545, 720, 398), (420, 398, 275, 398)])
    erd = DIAGRAMS / "erd.png"
    diagram_canvas(erd, "CyberShield Logical Data Model", [
        (80, 285, 290, 185, "USER\nuid (PK)\nemail\nprofile fields"),
        (550, 100, 300, 210, "ANALYSIS\nid (PK)\ncontent\nriskScore\nverdict\nanalyzedAt"),
        (550, 520, 300, 160, "RESULT DETAIL\nindicators\nfeatures\nscreenshotUrl"),
        (1030, 315, 270, 140, "SETTINGS\nreport format\npreferences"),
    ], [(370, 340, 550, 205), (370, 420, 550, 600), (850, 205, 1030, 360), (700, 310, 700, 520)])
    return use_case, dfd, erd


def title_page(doc):
    p(doc, "", before=95)
    p(doc, "CYBERSHIELD", WD_ALIGN_PARAGRAPH.CENTER, 20, bold=True, after=8)
    p(doc, "Software Requirements Specification", WD_ALIGN_PARAGRAPH.CENTER, 18, bold=True, after=32)
    p(doc, "An Evidence-Driven URL Phishing Detection Platform", WD_ALIGN_PARAGRAPH.CENTER, 14, italic=True, after=42)
    p(doc, "LDRP Institute of Technology and Research", WD_ALIGN_PARAGRAPH.CENTER, 14, bold=True, after=8)
    p(doc, "Academic Project Report", WD_ALIGN_PARAGRAPH.CENTER, 13, after=32)
    p(doc, "Submitted by: Neel Savsani", WD_ALIGN_PARAGRAPH.CENTER, 12, after=5)
    p(doc, "Project Guide: ______________________________", WD_ALIGN_PARAGRAPH.CENTER, 12, after=5)
    p(doc, f"Date: {date.today().strftime('%d %B %Y')}", WD_ALIGN_PARAGRAPH.CENTER, 12)
    doc.add_page_break()


def front(doc):
    heading(doc, "Acknowledgement", 1)
    p(doc, "This Software Requirements Specification has been prepared for the CyberShield academic project. The report documents the requirements and implementation completed to date. We acknowledge the guidance of faculty members, project reviewers, peers, and the open-source communities whose tools support the design, implementation, testing, and documentation of the system.")
    doc.add_page_break()
    heading(doc, "Abstract", 1)
    p(doc, "CyberShield is a web-based URL phishing analysis platform that investigates a submitted public HTTP or HTTPS address using live infrastructure and page-behaviour evidence. The system performs URL safety checks, HTTP, DNS, domain, TLS, browser, HTML, JavaScript, form, and optional reputation analysis. It produces a risk probability, a user-facing verdict, contributing indicators, a feature breakdown, and a scan-specific screenshot when capture succeeds. The current classifier is an explainable baseline intended for decision support; it is not a validated automated blocking model. This document specifies the system using the IEEE 830-oriented SRS structure and records the implemented functionality, constraints, data handling, interfaces, and planned extensions.")
    doc.add_page_break()
    heading(doc, "Table of Contents", 1)
    for item in ["1. Introduction", "2. Overall Description", "3. Specific Requirements", "4. System Models", "5. Appendices"]:
        p(doc, item, WD_ALIGN_PARAGRAPH.LEFT, 12, after=3)
    doc.add_page_break()


def introduction(doc):
    heading(doc, "1. Introduction", 1)
    heading(doc, "1.1 Purpose", 2)
    p(doc, "CyberShield assists users in assessing the phishing risk of a URL before they interact with it. The system collects evidence from multiple sources rather than relying only on URL spelling, brand keywords, or fixed character-count rules. It is designed to explain the observed evidence so that the user can make an informed decision.")
    heading(doc, "1.2 Scope", 2)
    p(doc, "The implemented scope is public URL analysis. A user can submit an HTTP or HTTPS URL from the dashboard, receive an analysis result, save it to personal history, reopen it from History or Reports, and print the result as a PDF. Email/text analysis and image upload are visible interface options but are not currently implemented as analysis workflows.")
    heading(doc, "1.3 Definitions, Acronyms and Abbreviations", 2)
    table(doc, ["Term", "Definition"], [
        ("SRS", "Software Requirements Specification."), ("SSRF", "Server-Side Request Forgery; a risk in which a submitted URL attempts to reach internal services."),
        ("TLS", "Transport Layer Security, used by HTTPS."), ("DNS", "Domain Name System."),
        ("Evidence-v1", "CyberShield feature schema derived from collected evidence."),
    ], [2200, 6800], "Table 1: Definitions and Abbreviations")
    heading(doc, "1.4 References", 2)
    for text in ["ISO/IEC/IEEE 29148 - Requirements engineering.", "FastAPI, Playwright Python, Firebase, and OWASP SSRF Prevention documentation.", "CyberShield source code, API models, frontend scripts, and current project README."]:
        bullet(doc, text)
    heading(doc, "1.5 Document Overview", 2)
    p(doc, "Section 2 describes the product and operating context. Section 3 specifies functional, non-functional, interface, and performance requirements. Section 4 provides logical system models. Section 5 includes traceability and implementation-status appendices.")


def overall(doc):
    heading(doc, "2. Overall Description", 1)
    heading(doc, "2.1 Product Perspective", 2)
    p(doc, "CyberShield is a multi-page web application backed by a FastAPI service. The frontend provides dashboard, authentication, history, reports, profile, settings, and result views. The backend validates URLs, coordinates collectors, converts evidence into stable features, and returns an explainable classification. Firebase Authentication and Firestore support the frontend account and saved-analysis experience, while SQLite-backed platform endpoints are available for local API workflows.")
    heading(doc, "2.2 Product Functions", 2)
    for text in ["Accept and normalize a public HTTP/HTTPS URL.", "Block localhost and private or non-public network destinations before collection.", "Collect HTTP, DNS, domain, TLS, browser, HTML, JavaScript, form, and optional reputation evidence.", "Capture a screenshot from the controlled browser session when available.", "Calculate an explainable risk score and show indicators, features, and user guidance.", "Store and reopen a registered user's complete analysis record."]:
        bullet(doc, text)
    heading(doc, "2.3 User Classes and Characteristics", 2)
    table(doc, ["User Class", "Characteristics"], [
        ("Guest", "Can submit URLs and use a temporary dashboard history; account-only history, reports, profile, and settings are restricted."),
        ("Registered user", "Uses Firebase Authentication and can retain, view, filter, delete, print, and revisit personal saved analyses."),
        ("Administrator/developer", "Runs the API, maintains configuration, configures optional reputation providers, and monitors the project environment."),
    ], [2300, 6700], "Table 2: User Classes")
    heading(doc, "2.4 Operating Environment", 2)
    p(doc, "The frontend runs in modern desktop or mobile browsers. The backend runs in a Python environment supporting FastAPI and Playwright Chromium. Network access is required for public targets, DNS and TLS queries, optional threat-intelligence providers, and Firebase cloud services. On Windows, browser-capable analysis uses a Proactor event loop in a worker thread to support Playwright subprocess creation.")
    heading(doc, "2.5 Constraints, Assumptions and Dependencies", 2)
    p(doc, "The system depends on public network availability, target-site behaviour, browser compatibility, and external services. A missing provider key or failed collector must be represented as unavailable evidence, not as a false safe or malicious result. The risk score is advisory: it must not be used as an automatic production blocking decision until a versioned model has been trained and validated on labelled data.")


def specific(doc):
    heading(doc, "3. Specific Requirements", 1)
    heading(doc, "3.1 Functional Requirements", 2)
    table(doc, ["ID", "Requirement"], [
        ("FR-01", "The system shall provide registration, sign-in, password recovery, sign-out, and guest entry."),
        ("FR-02", "The system shall accept complete HTTP/HTTPS URLs and reject malformed, local, private, link-local, reserved, or otherwise non-public targets before analysis."),
        ("FR-03", "The system shall collect HTTP reachability/redirect, DNS, domain age, TLS certificate, and configured reputation evidence when available."),
        ("FR-04", "The system shall render an admitted URL in a controlled browser context and collect final URL, title, browser events, page content, forms, and a scan-specific screenshot when possible."),
        ("FR-05", "The system shall derive evidence-v1 features and return a phishing probability, risk level, ranked contributors, and model status notice."),
        ("FR-06", "The Result page shall show score, verdict, input, timestamp, recommendation, indicators, feature values, and the matching screenshot."),
        ("FR-07", "Registered users' saved analysis records shall retain content, score, verdict, indicators, features, screenshot URL, and timestamps."),
        ("FR-08", "History and Reports shall reconstruct the same complete result data as the initial analysis result when a user selects View."),
        ("FR-09", "The system shall allow a saved result to be printed or saved as a PDF through the browser print flow."),
    ], [1100, 7900], "Table 3: Functional Requirements")
    heading(doc, "3.2 Non-functional Requirements", 2)
    table(doc, ["ID", "Category", "Requirement"], [
        ("NFR-01", "Security", "Reject unsafe destinations before any HTTP, TLS, or browser collection to reduce SSRF risk."),
        ("NFR-02", "Privacy", "A screenshot must belong to the particular analysis that generated it; the Result page must not fall back to a shared latest screenshot."),
        ("NFR-03", "Reliability", "Collector failures shall return partial labelled evidence and shall not prevent browser resources from closing."),
        ("NFR-04", "Usability", "The interface shall explain risk using a concise verdict, contributors, recommendation, and accessible feature values."),
        ("NFR-05", "Maintainability", "Collectors, feature extraction, classification, and frontend rendering shall remain separated components."),
        ("NFR-06", "Ethical use", "The baseline/synthetic model limitation shall be documented before any automated blocking is considered."),
    ], [1100, 1700, 6200], "Table 4: Non-functional Requirements")
    heading(doc, "3.3 External Interface Requirements", 2)
    heading(doc, "3.3.1 User Interface", 3)
    p(doc, "The Dashboard shall provide URL input, progress, summary statistics, and recent scans. The Result page shall display complete result evidence. History and Reports shall allow a user to find and reopen saved analyses. Profile and Settings are account-only screens.")
    heading(doc, "3.3.2 API Interface", 3)
    table(doc, ["Endpoint", "Method", "Purpose"], [
        ("/analyze", "POST", "Returns the synchronous evidence-driven URL analysis response."),
        ("/api/scans", "POST/GET", "Submits asynchronous platform scans and retrieves local platform scan records."),
        ("/api/scans/{scan_id}", "GET", "Retrieves a specific queued scan after access checks."),
        ("/health", "GET", "Reports that the backend service is online."),
        ("/reports", "Static", "Serves scan-specific browser screenshot files."),
    ], [2500, 1400, 5100], "Table 5: API Interfaces")
    heading(doc, "3.4 Performance Requirements", 2)
    p(doc, "The system shall promptly reject invalid or unsafe targets. For admitted URLs, execution time is dependent on live web, DNS, TLS, browser, and provider response conditions; therefore the interface shall show progress while analysis is active. Browser analysis shall execute away from the server's main event loop where the platform requires it. External network calls shall use bounded timeouts and failures shall produce a readable result state.")
    heading(doc, "3.5 Design Constraints", 2)
    p(doc, "CyberShield currently supports only publicly reachable HTTP/HTTPS URLs. Optional VirusTotal and Google Safe Browsing checks require user-supplied server-side API keys. OpenPhish and URLhaus are declared future feed integrations. The deployed production environment must apply restrictive CORS, secret management, HTTPS, Firebase Security Rules, rate limiting, and retained-data controls.")


def models(doc, diagrams):
    heading(doc, "4. System Models", 1)
    heading(doc, "4.1 Use Case Model", 2)
    p(doc, "The main actors are guest users, registered users, and administrators. The use case model represents the implemented user workflow and the administrative operational role.")
    doc.add_picture(str(diagrams[0]), width=Inches(5.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 1: CyberShield Use Case Model", above=False)
    heading(doc, "4.2 Data Flow Diagram", 2)
    p(doc, "At level 0, the user submits a URL to CyberShield. The platform exchanges requests with public web/DNS/TLS services and optional reputation providers, then stores the user-visible saved result in Firestore for authenticated users.")
    doc.add_picture(str(diagrams[1]), width=Inches(5.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 2: CyberShield Level-0 Data Flow Diagram", above=False)
    heading(doc, "4.3 Logical Data Model", 2)
    p(doc, "The application maintains user-owned analyses. Each analysis retains the submitted content, verdict, risk score, timestamp, indicators, extracted features, and screenshot URL. User settings are associated with the same authenticated user scope.")
    doc.add_picture(str(diagrams[2]), width=Inches(5.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 3: CyberShield Logical Data Model", above=False)
    heading(doc, "4.4 Analysis Sequence", 2)
    for item in ["Normalize and validate the URL.", "Resolve hostname and reject unsafe/non-public targets.", "Collect infrastructure and reputation evidence.", "Render the page in Playwright and inspect page, script, and form evidence.", "Extract evidence-v1 features and predict risk with contributor explanations.", "Save the complete user-visible record and render the result."]:
        number(doc, item)


def appendices(doc):
    heading(doc, "5. Appendices", 1)
    heading(doc, "Appendix A - Requirement Traceability Matrix", 2)
    table(doc, ["Requirement", "Implementation evidence", "Validation"], [
        ("FR-02", "URL safety gate and orchestrator", "Submit malformed, localhost, and private targets; confirm rejection before collection."),
        ("FR-03/FR-04", "Analyzer modules and Playwright browser session", "Analyze an admitted public target and inspect returned evidence/screenshot."),
        ("FR-05/FR-06", "Feature extractor, classifier, result.js", "Verify probability, contributors, indicators, features, and verdict display."),
        ("FR-07/FR-08", "Firebase data layer, history.js, reports.js", "Save a result and reopen it from History and Reports; compare detail with the fresh result."),
        ("NFR-02", "Per-analysis screenshot URL mapping", "Confirm a saved result shows only its own screenshot or hides the panel when absent."),
    ], [1500, 4200, 3300], "Table 6: Requirement Traceability Matrix")
    heading(doc, "Appendix B - Current Implementation Status", 2)
    table(doc, ["Area", "Status", "Notes"], [
        ("URL analysis", "Implemented", "Public URL analysis via live collector pipeline."),
        ("Browser screenshot", "Implemented", "Stored and restored as part of the matching analysis record."),
        ("History/Reports result fidelity", "Implemented", "Features and screenshot URL are passed into the Result page."),
        ("Email/text and image analysis", "Planned", "Interface placeholders exist; analysis workflow is not implemented."),
        ("OpenPhish/URLhaus integration", "Planned", "Provider placeholders are present; production integration remains future work."),
        ("Validated production model", "Planned", "Baseline/demonstration model must be replaced by a versioned validated model."),
    ], [2500, 1500, 5000], "Table 7: Current Implementation Status")


def main():
    diagrams = make_diagrams()
    doc = Document()
    setup(doc)
    title_page(doc)
    front(doc)
    introduction(doc)
    overall(doc)
    specific(doc)
    models(doc, diagrams)
    appendices(doc)
    doc.core_properties.title = "CyberShield Software Requirements Specification"
    doc.core_properties.author = "CyberShield Project Team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
